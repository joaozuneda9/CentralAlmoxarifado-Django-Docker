"""Leitura de .xlsx (sem ORM, sem regras de negócio), usada pela carga do de/para.

`ler_xlsx` devolve as linhas de cada aba pedida como dicionários cabeçalho → valor da célula; a
conversão dos valores (`texto_celula`) fica para quem usa a leitura, que sabe o que cada coluna
deveria conter.
"""

from pathlib import Path

import openpyxl

Linha = dict[str, object]


class ErroPlanilha(ValueError):
    """A planilha não tem a forma esperada (aba ou coluna faltando)."""


def texto_celula(valor: object) -> str:
    """Texto da célula: `None` → "", inteiro sem casas, sem espaços nas pontas nem `'` inicial."""
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    if isinstance(valor, str):
        texto = valor.strip()
        # Planilhas guardam com um apóstrofo na frente o texto que parece fórmula.
        return texto[1:] if texto.startswith("'") else texto
    return str(valor)


def _vazia(valor: object) -> bool:
    return valor is None or (isinstance(valor, str) and valor == "")


def ler_xlsx(caminho: Path, abas: dict[str, tuple[str, ...]]) -> dict[str, list[Linha]]:
    """Lê as abas pedidas, e só elas, com as colunas obrigatórias de cada uma. Cada linha leva o
    número real na chave `_linha`."""
    pedidas = abas
    livro = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    try:
        problemas = [
            f'A planilha não tem a aba "{nome}".'
            for nome in pedidas
            if nome not in livro.sheetnames
        ]
        abas: dict[str, list[Linha]] = {}
        for nome, obrigatorias in pedidas.items():
            if nome not in livro.sheetnames:
                continue
            linhas = livro[nome].iter_rows(values_only=True)
            cabecalho = next(linhas, ())
            colunas: dict[str, int] = {}
            for posicao, celula in enumerate(cabecalho):
                titulo = texto_celula(celula)
                if titulo and titulo not in colunas:
                    colunas[titulo] = posicao
            problemas += [
                f'{nome}: coluna "{coluna}" não encontrada.'
                for coluna in obrigatorias
                if coluna not in colunas
            ]
            abas[nome] = [
                {
                    **{
                        titulo: (celulas[posicao] if posicao < len(celulas) else None)
                        for titulo, posicao in colunas.items()
                    },
                    "_linha": numero,
                }
                for numero, celulas in enumerate(linhas, start=2)
                if not all(_vazia(celula) for celula in celulas)
            ]
    finally:
        livro.close()
    if problemas:
        raise ErroPlanilha("\n".join(problemas))
    return abas
